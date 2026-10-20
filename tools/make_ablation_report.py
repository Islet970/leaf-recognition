"""生成消融实验图表与中文 Word 报告（plans.md §4.2 / §4.4 / §14.6，T6.2）。

消融组固定为 A（基线）/ E（+CBAM）/ C（+MSFF）/ D（MSFF+CBAM），B 不参与。
数据源全部为已有产物，只读取、不修改：

    五类白底 test（58）   runs/spec5/{A,E,C,D}*_seed42/test_metrics.json
    五类实景 test（52）   runs/live5/{A,E,C,D}*_seed42/test_metrics.json
    32 类白底 test（358） runs/spec32/{A,E,C,D}*_seed42/test_metrics.json
    困难帧 R1/R2/R3（各 227） pic/melt-test/raw/metrics_robust_{w5,l5,w32}.csv

输出（默认 pic/melt-test/）：
    ablation_metrics.csv                 6 数据集 × 4 模型
    ablation_main_accuracy.png           主测试分组柱状图
    ablation_robustness_accuracy.png     困难帧分组柱状图
    ablation_gains.png                   模块增益（百分点），左主测试 / 右困难帧
    消融实验报告.docx                    简短中文报告（python-docx）

示例：
    python tools/make_ablation_report.py
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PROJECT_ROOT = Path(__file__).resolve().parents[1]

# 与 evaluate_live.py / make_comparison_plots.py 保持同一配色：A 蓝、E 紫、C 绿、D 红
COLORS = {"A": "#1f77b4", "E": "#9467bd", "C": "#2ca02c", "D": "#d62728"}
GAIN_COLORS = {"CBAM": "#9467bd", "MSFF": "#2ca02c", "组合": "#d62728"}
MODELS = ["A", "E", "C", "D"]
MODEL_LABELS = {
    "A": "A 基线",
    "E": "E +CBAM",
    "C": "C +MSFF",
    "D": "D MSFF+CBAM",
}
MODEL_DIRS = {
    "A": "A_mnv3s",
    "E": "E_mnv3s_cbam",
    "C": "C_leafnet_msff",
    "D": "D_leafnet_msff_cbam",
}

# 六个数据集：key, 图中标签, 来源说明
DATASETS = [
    {"key": "spec5", "label": "五类白底", "full": "五类白底 test", "kind": "main"},
    {"key": "live5", "label": "五类实景", "full": "五类实景 test", "kind": "main"},
    {"key": "spec32", "label": "32 类白底", "full": "32 类白底 test", "kind": "main"},
    {"key": "robust_w5", "label": "R1（W5）", "full": "R1：W5 → 困难帧", "kind": "robust"},
    {"key": "robust_l5", "label": "R2（L5）", "full": "R2：L5 → 困难帧", "kind": "robust"},
    {"key": "robust_w32", "label": "R3（W32）", "full": "R3：W32 → 困难帧", "kind": "robust"},
]
MAIN_KEYS = [d["key"] for d in DATASETS if d["kind"] == "main"]
ROBUST_KEYS = [d["key"] for d in DATASETS if d["kind"] == "robust"]
DATASET_LABEL = {d["key"]: d["label"] for d in DATASETS}

RUN_DIRS = {"spec5": "runs/spec5", "live5": "runs/live5", "spec32": "runs/spec32"}
ROBUST_CSVS = {
    "robust_w5": "pic/melt-test/raw/metrics_robust_w5.csv",
    "robust_l5": "pic/melt-test/raw/metrics_robust_l5.csv",
    "robust_w32": "pic/melt-test/raw/metrics_robust_w32.csv",
}


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="汇总 A/E/C/D 消融指标，生成柱状图与中文 Word 报告。"
    )
    parser.add_argument(
        "--out",
        default="pic/melt-test",
        help="输出目录（默认 pic/melt-test，相对路径按项目根目录解析）",
    )
    parser.add_argument("--seed", type=int, default=42, help="run 目录 seed 后缀（默认 42）")
    parser.add_argument("--dpi", type=int, default=180, help="PNG DPI（默认 180）")
    return parser.parse_args(argv)


def project_path(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else PROJECT_ROOT / path


def configure_fonts() -> None:
    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei", "Noto Sans SC"]
    plt.rcParams["font.family"] = "sans-serif"
    plt.rcParams["axes.unicode_minus"] = False


def load_json_metrics(path: Path, model: str, dataset: str) -> tuple[dict, dict]:
    if not path.is_file():
        raise SystemExit(f"缺少指标文件: {path}")
    try:
        obj = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SystemExit(f"无法读取 {path}: {exc}") from exc
    if "model_id" not in obj:
        print(f"[警告] {path} 缺少 model_id，按目标模型 {model} 处理")
    elif str(obj["model_id"]).upper() != model:
        raise SystemExit(
            f"{path} 的 model_id={obj.get('model_id')!r} 与目标模型 {model!r} 不一致"
        )
    for field in ("accuracy", "macro_f1", "ce_loss", "num_samples"):
        if field not in obj:
            raise SystemExit(f"{path} 缺少字段 {field!r}")
    top5 = obj.get("top5_accuracy", obj.get("topk_accuracy"))
    extra = {
        "top5_accuracy": float(top5) if top5 is not None else None,
        "checkpoint_epoch": (
            int(obj["checkpoint_epoch"]) if "checkpoint_epoch" in obj else None
        ),
    }
    row = {
        "dataset": dataset,
        "model": model,
        "accuracy": float(obj["accuracy"]),
        "macro_f1": float(obj["macro_f1"]),
        "ce_loss": float(obj["ce_loss"]),
        "n": int(obj["num_samples"]),
    }
    return row, extra


def load_csv_metrics(path: Path, dataset: str) -> list[dict]:
    if not path.is_file():
        raise SystemExit(f"缺少指标文件: {path}")
    df = pd.read_csv(path)
    required = ["model", "accuracy", "macro_f1", "ce_loss", "n"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise SystemExit(f"{path} 缺少列 {missing}")
    df["model"] = df["model"].astype(str).str.strip().str.upper()
    rows: list[dict] = []
    for model in MODELS:
        sub = df[df["model"] == model]
        if len(sub) != 1:
            raise SystemExit(f"{path} 中模型 {model} 应有且仅有 1 行，实际 {len(sub)} 行")
        row = sub.iloc[0]
        rows.append(
            {
                "dataset": dataset,
                "model": model,
                "accuracy": float(row["accuracy"]),
                "macro_f1": float(row["macro_f1"]),
                "ce_loss": float(row["ce_loss"]),
                "n": int(row["n"]),
            }
        )
    return rows


def collect_metrics(seed: int) -> tuple[pd.DataFrame, dict[tuple[str, str], dict]]:
    rows: list[dict] = []
    extras: dict[tuple[str, str], dict] = {}
    for dataset, run_dir in RUN_DIRS.items():
        for model in MODELS:
            path = (
                PROJECT_ROOT
                / run_dir
                / f"{MODEL_DIRS[model]}_seed{seed}"
                / "test_metrics.json"
            )
            row, extra = load_json_metrics(path, model, dataset)
            rows.append(row)
            extras[(dataset, model)] = extra
    for dataset, csv_path in ROBUST_CSVS.items():
        rows.extend(load_csv_metrics(PROJECT_ROOT / csv_path, dataset))
    df = pd.DataFrame(rows, columns=["dataset", "model", "accuracy", "macro_f1", "ce_loss", "n"])

    # 同一数据集的四模型必须来自同一份测试清单，样本数应完全一致
    for dataset, group in df.groupby("dataset", sort=False):
        counts = group["n"].unique()
        if len(counts) != 1:
            raise SystemExit(
                f"数据集 {dataset} 的四个模型 n 不一致: "
                f"{group.set_index('model')['n'].to_dict()}"
            )
        models = sorted(group["model"].tolist())
        if models != sorted(MODELS):
            raise SystemExit(f"数据集 {dataset} 的模型集合不完整: {models}")
    robust_counts = {int(df.loc[df["dataset"] == k, "n"].iloc[0]) for k in ROBUST_KEYS}
    if len(robust_counts) != 1:
        raise SystemExit(f"三组困难帧样本数不一致: {sorted(robust_counts)}")
    return df, extras


def pivot_accuracy(df: pd.DataFrame) -> pd.DataFrame:
    table = df.pivot(index="dataset", columns="model", values="accuracy")
    return table.reindex(index=[d["key"] for d in DATASETS], columns=MODELS)


def _style_axes(ax, ylabel: str, title: str) -> None:
    ax.set_ylabel(ylabel, fontsize=11)
    ax.set_title(title, fontsize=13, fontweight="bold")
    ax.grid(True, axis="y", linestyle="--", alpha=0.35)
    ax.set_axisbelow(True)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)


def plot_grouped_accuracy(
    df: pd.DataFrame,
    keys: list[str],
    out_path: Path,
    title: str,
    ylim: float,
    dpi: int,
    baseline: dict | None = None,
) -> None:
    """分组柱状图：每组四个模型，柱顶标百分比。

    baseline 可选，形式为 {"key": 数据集 key, "value": 基线值, "label": 文字}；
    只在指定数据集那一组画一段短虚线，避免被误读为其他数据集的基线。
    """
    configure_fonts()
    acc = df.set_index(["dataset", "model"])["accuracy"]
    ns = df.groupby("dataset")["n"].first()

    x = np.arange(len(keys))
    width = 0.78 / len(MODELS)
    fig, ax = plt.subplots(figsize=(9.2, 5.0))
    baseline_value = float(baseline["value"]) if baseline is not None else None
    for i, model in enumerate(MODELS):
        values = [float(acc.loc[(k, model)]) for k in keys]
        positions = x + (i - (len(MODELS) - 1) / 2) * width
        bars = ax.bar(
            positions,
            values,
            width=width,
            color=COLORS[model],
            edgecolor="black",
            linewidth=0.6,
            label=MODEL_LABELS[model],
        )
        for idx, (bar, value) in enumerate(zip(bars, values)):
            label_y = value + 0.012
            if (
                baseline is not None
                and baseline.get("key") == keys[idx]
                and value <= baseline_value
            ):
                # 该组柱子接近基线，把柱顶标签抬到虚线上方，避免与基线文字重叠
                label_y = baseline_value + 0.045
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                min(label_y, ylim),
                f"{value * 100:.2f}%",
                ha="center",
                va="bottom",
                fontsize=8.2,
            )
    ax.set_xticks(x)
    ax.set_xticklabels(
        [f"{DATASET_LABEL[k]}\n(n={int(ns.loc[k])})" for k in keys], fontsize=10.5
    )
    ax.set_ylim(0.0, ylim)
    ax.set_xlim(-0.6, len(keys) - 0.4)
    if baseline is not None and baseline.get("key") in keys:
        pos = keys.index(baseline["key"])
        value = float(baseline["value"])
        ax.hlines(
            value,
            pos - 0.48,
            pos + 0.48,
            colors="#555555",
            linestyles=(0, (4, 3)),
            linewidth=1.2,
            zorder=5,
        )
        ax.text(
            pos,
            value + ylim * 0.11,
            baseline["label"],
            ha="center",
            va="bottom",
            fontsize=8.5,
            color="#333333",
        )
    _style_axes(ax, "准确率", title)
    ax.legend(ncol=4, fontsize=9, frameon=False, loc="upper center", bbox_to_anchor=(0.5, -0.16))
    fig.tight_layout()
    fig.savefig(out_path, dpi=dpi, bbox_inches="tight")
    plt.close(fig)


def gains_table(df: pd.DataFrame) -> pd.DataFrame:
    """返回每个数据集的三项增益（百分点）：CBAM=E−A，MSFF=C−A，组合=D−A。"""
    acc = df.set_index(["dataset", "model"])["accuracy"]
    rows = []
    for spec in DATASETS:
        key = spec["key"]
        a = float(acc.loc[(key, "A")])
        rows.append(
            {
                "dataset": key,
                "CBAM": (float(acc.loc[(key, "E")]) - a) * 100,
                "MSFF": (float(acc.loc[(key, "C")]) - a) * 100,
                "组合": (float(acc.loc[(key, "D")]) - a) * 100,
            }
        )
    return pd.DataFrame(rows).set_index("dataset")


def plot_gains(
    df: pd.DataFrame, keys_panels: list[list[str]], out_path: Path, dpi: int
) -> None:
    """左右两个子图：主测试增益与困难帧增益，单位为百分点。"""
    configure_fonts()
    gains = gains_table(df)
    series = ["CBAM", "MSFF", "组合"]
    labels = {"CBAM": "CBAM (E−A)", "MSFF": "MSFF (C−A)", "组合": "组合 (D−A)"}
    titles = ["主测试增益", "困难帧增益"]

    fig, axes = plt.subplots(1, 2, figsize=(13.0, 4.9))
    for ax, keys, title in zip(axes, keys_panels, titles):
        values = gains.loc[keys, series].to_numpy(dtype=float)
        lo = min(0.0, float(values.min()))
        hi = max(0.0, float(values.max()))
        span = max(hi - lo, 1.0)
        y_low, y_high = lo - span * 0.22, hi + span * 0.22
        x = np.arange(len(keys))
        width = 0.72 / len(series)
        for i, name in enumerate(series):
            vals = values[:, i]
            positions = x + (i - (len(series) - 1) / 2) * width
            bars = ax.bar(
                positions,
                vals,
                width=width,
                color=GAIN_COLORS[name],
                edgecolor="black",
                linewidth=0.6,
                label=labels[name],
            )
            for bar, value in zip(bars, vals):
                if value >= 0:
                    text_y, va = value + span * 0.025, "bottom"
                else:
                    text_y, va = value - span * 0.025, "top"
                ax.text(
                    bar.get_x() + bar.get_width() / 2,
                    text_y,
                    f"{value:+.2f}",
                    ha="center",
                    va=va,
                    fontsize=8.5,
                )
        ax.axhline(0.0, color="black", linewidth=0.9)
        ax.set_xticks(x)
        ax.set_xticklabels([DATASET_LABEL[k] for k in keys], fontsize=10.5)
        ax.set_xlim(-0.55, len(keys) - 0.45)
        ax.set_ylim(y_low, y_high)
        _style_axes(ax, "准确率增益（百分点）", title)
        ax.legend(fontsize=8.5, frameon=False, loc="upper left")

    fig.tight_layout()
    fig.savefig(out_path, dpi=dpi, bbox_inches="tight")
    plt.close(fig)


# --------------------------------------------------------------------------- #
# Word 报告
# --------------------------------------------------------------------------- #

BODY_FONT = "宋体"
HEAD_FONT = "微软雅黑"


def set_run_font(
    run,
    size: float = 10.5,
    bold: bool = False,
    east: str = BODY_FONT,
    latin: str = "Times New Roman",
) -> None:
    """同时设置西文与东亚字体（w:eastAsia），避免中文回退到默认字体。"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def add_para(
    doc: Document,
    text: str,
    size: float = 10.5,
    bold: bool = False,
    east: str = BODY_FONT,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    indent: bool = False,
    space_after: float = 3.0,
    space_before: float = 0.0,
) -> None:
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.25
    if indent:
        pf.first_line_indent = Pt(size * 2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, east=east, latin="Microsoft YaHei" if east == HEAD_FONT else "Times New Roman")


def add_heading(doc: Document, text: str) -> None:
    add_para(
        doc,
        text,
        size=13,
        bold=True,
        east=HEAD_FONT,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=8,
        space_after=4,
    )


def add_caption(doc: Document, text: str) -> None:
    add_para(
        doc,
        text,
        size=9,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing = 1.25
    run = p.add_run(text)
    set_run_font(run, size=10.5)


def shade_cell(cell, fill: str = "D9E2F3") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 10.0) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_accuracy_table(doc: Document, df: pd.DataFrame) -> None:
    acc = pivot_accuracy(df)
    ns = df.groupby("dataset")["n"].first()
    table = doc.add_table(rows=1 + len(DATASETS), cols=5)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Cm(5.6), Cm(2.55), Cm(2.55), Cm(2.55), Cm(2.55)]

    headers = ["测试集", "A 基线", "E +CBAM", "C +MSFF", "D 组合"]
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_text(cell, header, bold=True)
        shade_cell(cell)
        cell.width = widths[j]
    for i, spec in enumerate(DATASETS, start=1):
        key = spec["key"]
        row = table.rows[i]
        set_cell_text(row.cells[0], f"{spec['full']}（n={int(ns.loc[key])}）", size=9.5)
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        row.cells[0].width = widths[0]
        for j, model in enumerate(MODELS, start=1):
            set_cell_text(row.cells[j], f"{float(acc.loc[key, model]) * 100:.2f}")
            row.cells[j].width = widths[j]


def format_model_values(
    values: dict[str, float | None],
    scale: float = 100.0,
    digits: int = 2,
    unit: str = "%",
) -> str:
    """把四模型数值格式化为中文描述；相等时合并，不相等时逐模型列出。"""
    clean = {m: values.get(m) for m in MODELS}
    if any(v is None or not np.isfinite(float(v)) for v in clean.values()):
        items = []
        for m in MODELS:
            v = clean[m]
            if v is None or not np.isfinite(float(v)):
                items.append(f"{m} 未提供")
            else:
                items.append(f"{m} {float(v) * scale:.{digits}f}{unit}")
        return "、".join(items)

    rounded = {m: round(float(clean[m]) * scale, digits) for m in MODELS}
    groups: dict[float, list[str]] = {}
    for m in MODELS:
        groups.setdefault(rounded[m], []).append(m)
    if len(groups) == 1:
        return f"A/E/C/D 均为 {next(iter(groups)):.{digits}f}{unit}"
    return "，".join(
        f"{'/'.join(models)} 为 {value:.{digits}f}{unit}"
        for value, models in groups.items()
    )


def build_docx(
    df: pd.DataFrame,
    extras: dict[tuple[str, str], dict],
    out_path: Path,
    main_png: Path,
    robust_png: Path,
    seed: int,
) -> None:
    acc = pivot_accuracy(df)
    ns = df.groupby("dataset")["n"].first()

    def pct(key: str, model: str) -> float:
        return float(acc.loc[key, model]) * 100

    def values_of(key: str) -> dict[str, float]:
        return {m: float(acc.loc[key, m]) for m in MODELS}

    def interaction(key: str) -> float:
        """交互项 (D−C)−(E−A)，单位为百分点。"""
        return (pct(key, "D") - pct(key, "C")) - (pct(key, "E") - pct(key, "A"))

    n_spec5 = int(ns["spec5"])
    n_live5 = int(ns["live5"])
    n_spec32 = int(ns["spec32"])
    n_robust = int(ns["robust_w5"])
    live_pp = 100.0 / n_live5
    live_diffs = [
        abs(pct("live5", "C") - pct("live5", "A")),
        abs(pct("live5", "D") - pct("live5", "A")),
        abs(pct("live5", "E") - pct("live5", "A")),
    ]
    sample_lo = min(live_diffs) / live_pp
    sample_hi = max(live_diffs) / live_pp

    inter_live = interaction("live5")
    inter_r1 = interaction("robust_w5")
    inter_r2 = interaction("robust_l5")
    inter_r3 = interaction("robust_w32")

    live_d_epoch = extras.get(("live5", "D"), {}).get("checkpoint_epoch")
    live_d_epoch_text = f"best epoch={live_d_epoch}" if live_d_epoch is not None else "best epoch 未记录"

    spec32_top5 = {
        m: extras.get(("spec32", m), {}).get("top5_accuracy") for m in MODELS
    }

    doc = Document()

    # 页面与正文样式
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(3)

    # 标题
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("消融实验报告")
    set_run_font(run, size=18, bold=True, east=HEAD_FONT, latin="Microsoft YaHei")

    add_para(
        doc,
        f"模型：A（基线）/ E（+CBAM）/ C（+MSFF）/ D（MSFF+CBAM）；seed {seed}",
        size=9.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )

    add_heading(doc, "1  实验目的")
    add_para(
        doc,
        "在相同数据与训练配置下，分别加入 CBAM 注意力模块与 MSFF 多尺度特征融合模块，"
        "考察单一模块及两者组合对识别准确率的影响，重点判断架构差异在何种测试条件下"
        "才能被观测到。",
        indent=True,
    )

    add_heading(doc, "2  实验设置")
    add_para(
        doc,
        "消融组固定 A（MobileNetV3-S 基线）、E（+CBAM）、C（+MSFF）、D（MSFF+CBAM），"
        f"B 不参与。测试集共 6 套：五类白底 test（{n_spec5} 张）、五类实景 test"
        f"（{n_live5} 张）、32 类白底 test（{n_spec32} 张），以及三组各 {n_robust} 张"
        "困难帧压力测试（R1：W5 五类白底模型；R2：L5 五类实景模型；R3：W32 32 类白底模型）。"
        f"统一使用 seed {seed} 的 best checkpoint。报告以准确率为主要指标，"
        "macro-F1 与 CE 损失见 ablation_metrics.csv。",
        indent=True,
    )

    add_heading(doc, "3  结果")
    add_para(
        doc,
        "六套测试集上的准确率见表 1，主测试与困难帧对比分别见图 1、图 2；"
        "各模块相对基线的增益（百分点）另存为 pic/melt-test/ablation_gains.png。",
        indent=True,
    )
    add_caption(doc, "表 1  六套测试集准确率（%）")
    add_accuracy_table(doc, df)
    add_para(doc, "", size=3, space_after=0)

    doc.add_picture(str(main_png), width=Cm(12.9))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_caption(doc, "图 1  A/E/C/D 在三套主测试集上的准确率")

    doc.add_picture(str(robust_png), width=Cm(12.9))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    add_caption(
        doc,
        "图 2  A/E/C/D 在三组困难帧压力测试上的准确率"
        "（R3 接近 0；虚线为 32 类随机基线≈3.1%）",
    )

    add_heading(doc, "4  分析与结论")
    add_bullet(
        doc,
        f"五类白底 test 已饱和：{format_model_values(values_of('spec5'))}，"
        "架构差异被任务难度掩盖。",
    )
    add_bullet(
        doc,
        f"五类实景 test（n={n_live5}，1 张≈{live_pp:.1f} pp）差异有限："
        f"{format_model_values(values_of('live5'))}。C−A = "
        f"{pct('live5', 'C') - pct('live5', 'A'):+.2f} pp、D−A = "
        f"{pct('live5', 'D') - pct('live5', 'A'):+.2f} pp，约 {sample_lo:.0f}–{sample_hi:.0f} "
        f"张样本即可造成同等幅度，暂不足以判定稳定增益；交互项 (D−C)−(E−A) = "
        f"{inter_live:+.2f} pp。",
    )
    add_bullet(
        doc,
        f"32 类白底 test 准确率无差异：{format_model_values(values_of('spec32'))}；"
        f"top-5 准确率 {format_model_values(spec32_top5)}。CE 与 macro-F1 仍有差异"
        "（见 ablation_metrics.csv），但准确率层面未观察到架构差异。",
    )
    add_bullet(
        doc,
        f"困难帧结果方向不一致。R1（W5 模型 → {n_robust} 张困难帧）："
        f"{format_model_values(values_of('robust_w5'))}；D 为 {pct('robust_w5', 'D'):.2f}% 最高，"
        f"E {pct('robust_w5', 'E'):.2f}%、C {pct('robust_w5', 'C'):.2f}% 均未超过 A "
        f"{pct('robust_w5', 'A'):.2f}%。R2（L5 模型 → {n_robust} 张困难帧）："
        f"{format_model_values(values_of('robust_l5'))}；E−A = "
        f"{pct('robust_l5', 'E') - pct('robust_l5', 'A'):+.2f} pp，D 为 "
        f"{pct('robust_l5', 'D'):.2f}%（{live_d_epoch_text}）明显偏低。"
        f"R3（W32 模型 → {n_robust} 张困难帧）：{format_model_values(values_of('robust_w32'))}；"
        f"最高仅 {max(values_of('robust_w32').values()) * 100:.2f}%，低于 32 类随机基线"
        "（约 3.1%），不能据此评价架构。",
    )
    add_bullet(
        doc,
        "交互项 (D−C)−(E−A) 方向不一致："
        f"五类实景（live5）{inter_live:+.2f}、R1 {inter_r1:+.2f}、R2 {inter_r2:+.2f}、"
        f"R3 {inter_r3:+.2f} 个百分点，不能宣称 MSFF+CBAM 具有稳定互补增益。",
    )

    add_heading(doc, "5  局限")
    add_bullet(
        doc,
        f"测试集规模小（五类白底 {n_spec5} 张、实景 {n_live5} 张），单张样本约对应 "
        f"{100.0 / max(n_spec5, n_live5):.1f}–{100.0 / min(n_spec5, n_live5):.1f} "
        "个百分点，结果易受个别样本影响。",
    )
    add_bullet(doc, f"仅 seed {seed} 单次训练，无法判断模型间差异的稳定性。")
    add_bullet(doc, "困难帧来自与实景训练相同的 5 段视频，存在场景/拍摄条件重叠，R1–R3 只宜作鲁棒性压力测试。")
    add_bullet(doc, "R3 为 32 类模型在仅覆盖 5/32 类的困难帧上评测，指标偏低且不可与主测试直接比较。")

    doc.save(str(out_path))


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    out_dir = project_path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    df, extras = collect_metrics(args.seed)
    csv_path = out_dir / "ablation_metrics.csv"
    df.to_csv(csv_path, index=False, float_format="%.6f")
    print(f"指标 CSV: {csv_path}  ({len(df)} 行 × {len(df.columns)} 列)")

    acc = pivot_accuracy(df)
    print("准确率汇总（%）:")
    print((acc * 100).round(2).to_string())

    main_png = out_dir / "ablation_main_accuracy.png"
    robust_png = out_dir / "ablation_robustness_accuracy.png"
    gains_png = out_dir / "ablation_gains.png"

    plot_grouped_accuracy(
        df,
        MAIN_KEYS,
        main_png,
        title="主测试集准确率（A/E/C/D）",
        ylim=1.05,
        dpi=args.dpi,
    )
    print(f"主测试图: {main_png}")
    plot_grouped_accuracy(
        df,
        ROBUST_KEYS,
        robust_png,
        title="困难帧鲁棒性测试准确率（A/E/C/D）",
        ylim=1.0,
        dpi=args.dpi,
        baseline={
            "key": "robust_w32",
            "value": 1.0 / 32.0,
            "label": "32 类随机基线≈3.1%",
        },
    )
    print(f"鲁棒性图: {robust_png}")
    plot_gains(df, [MAIN_KEYS, ROBUST_KEYS], gains_png, dpi=args.dpi)
    print(f"增益图:   {gains_png}")

    print("增益（百分点）:")
    print(gains_table(df).round(2).to_string())

    docx_path = out_dir / "消融实验报告.docx"
    build_docx(df, extras, docx_path, main_png, robust_png, args.seed)
    print(f"Word 报告: {docx_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
